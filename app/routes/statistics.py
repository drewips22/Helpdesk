from datetime import datetime, timedelta
from io import BytesIO
from flask import Blueprint, render_template, request, send_file
from flask_login import login_required
from app.models import Ticket, TicketLog
from app.extensions import db
from app.utils.decorators import teknisi_required

stats_bp = Blueprint('statistics', __name__, url_prefix='/statistics')


@stats_bp.route('/')
@login_required
@teknisi_required
def index():
    # Time range
    range_filter = request.args.get('range', '30')
    try:
        days = int(range_filter)
    except ValueError:
        days = 30

    date_from = datetime.now() - timedelta(days=days)

    # Base query
    base_q = Ticket.query.filter(Ticket.created_at >= date_from)

    # Total stats
    total = base_q.count()
    open_count = base_q.filter_by(status='OPEN').count()
    assigned_count = base_q.filter_by(status='ASSIGNED').count()
    progress_count = base_q.filter_by(status='ON_PROGRESS').count()
    done_count = base_q.filter_by(status='DONE').count()
    closed_count = base_q.filter_by(status='CLOSED').count()

    # By urgency
    urgency_stats = {}
    for urg in ['rendah', 'sedang', 'tinggi', 'kritis']:
        urgency_stats[urg] = base_q.filter_by(urgency=urg).count()

    # By bagian (top 10)
    bagian_stats = db.session.query(
        Ticket.bagian, db.func.count(Ticket.id)
    ).filter(Ticket.created_at >= date_from)\
        .group_by(Ticket.bagian)\
        .order_by(db.func.count(Ticket.id).desc())\
        .limit(10).all()

    # By jenis_perangkat
    perangkat_stats = db.session.query(
        Ticket.jenis_perangkat, db.func.count(Ticket.id)
    ).filter(Ticket.created_at >= date_from)\
        .group_by(Ticket.jenis_perangkat)\
        .order_by(db.func.count(Ticket.id).desc())\
        .limit(10).all()

    # By jenis_masalah
    masalah_stats = db.session.query(
        Ticket.jenis_masalah, db.func.count(Ticket.id)
    ).filter(Ticket.created_at >= date_from)\
        .group_by(Ticket.jenis_masalah)\
        .order_by(db.func.count(Ticket.id).desc())\
        .limit(10).all()

    # Daily trend
    daily_trend = db.session.query(
        db.func.date(Ticket.created_at), db.func.count(Ticket.id)
    ).filter(Ticket.created_at >= date_from)\
        .group_by(db.func.date(Ticket.created_at))\
        .order_by(db.func.date(Ticket.created_at)).all()

    # Average resolution time (in hours)
    resolved_tickets = Ticket.query.filter(
        Ticket.resolved_at.isnot(None),
        Ticket.created_at >= date_from
    ).all()

    avg_resolution = 0
    if resolved_tickets:
        total_hours = sum(
            (t.resolved_at - t.created_at).total_seconds() / 3600
            for t in resolved_tickets
        )
        avg_resolution = round(total_hours / len(resolved_tickets), 1)

    # SLA compliance
    sla_compliant = 0
    sla_breached = 0
    for t in resolved_tickets:
        if t.sla_deadline and t.resolved_at <= t.sla_deadline:
            sla_compliant += 1
        else:
            sla_breached += 1

    sla_rate = round((sla_compliant / len(resolved_tickets) * 100), 1) if resolved_tickets else 0

    return render_template('statistics/index.html',
                           range_filter=range_filter,
                           total=total,
                           open_count=open_count,
                           assigned_count=assigned_count,
                           progress_count=progress_count,
                           done_count=done_count,
                           closed_count=closed_count,
                           urgency_stats=urgency_stats,
                           bagian_stats=bagian_stats,
                           perangkat_stats=perangkat_stats,
                           masalah_stats=masalah_stats,
                           daily_trend=daily_trend,
                           avg_resolution=avg_resolution,
                           sla_rate=sla_rate,
                           sla_compliant=sla_compliant,
                           sla_breached=sla_breached)


@stats_bp.route('/export')
@login_required
@teknisi_required
def export_excel():
    """Export tickets to Excel."""
    range_filter = request.args.get('range', '30')
    try:
        days = int(range_filter)
    except ValueError:
        days = 30

    date_from = datetime.now() - timedelta(days=days)

    tickets = Ticket.query.filter(
        Ticket.created_at >= date_from
    ).order_by(Ticket.created_at.desc()).all()

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        wb = Workbook()
        ws = wb.active
        ws.title = 'Laporan Tiket'

        # Title
        ws.merge_cells('A1:K1')
        title_cell = ws['A1']
        title_cell.value = f'Laporan Tiket SI-TIK - {days} Hari Terakhir'
        title_cell.font = Font(name='Arial', size=14, bold=True, color='1B2A4A')
        title_cell.alignment = Alignment(horizontal='center')

        ws.merge_cells('A2:K2')
        date_cell = ws['A2']
        date_cell.value = f'Periode: {date_from.strftime("%d/%m/%Y")} - {datetime.now().strftime("%d/%m/%Y")}'
        date_cell.font = Font(name='Arial', size=10, color='666666')
        date_cell.alignment = Alignment(horizontal='center')

        # Headers
        headers = ['No', 'Kode Tiket', 'Pelapor', 'Bagian', 'Lokasi',
                    'Perangkat', 'Masalah', 'Urgency', 'Status',
                    'Tanggal Dibuat', 'Tanggal Selesai']

        header_fill = PatternFill(start_color='1B2A4A', end_color='1B2A4A', fill_type='solid')
        header_font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = thin_border

        # Data
        for idx, ticket in enumerate(tickets, 1):
            row = idx + 4
            ws.cell(row=row, column=1, value=idx).border = thin_border
            ws.cell(row=row, column=2, value=ticket.ticket_code).border = thin_border
            ws.cell(row=row, column=3, value=ticket.pelapor.nama if ticket.pelapor else '-').border = thin_border
            ws.cell(row=row, column=4, value=ticket.bagian).border = thin_border
            ws.cell(row=row, column=5, value=ticket.lokasi).border = thin_border
            ws.cell(row=row, column=6, value=ticket.jenis_perangkat).border = thin_border
            ws.cell(row=row, column=7, value=ticket.jenis_masalah).border = thin_border
            ws.cell(row=row, column=8, value=ticket.urgency.capitalize()).border = thin_border
            ws.cell(row=row, column=9, value=ticket.status_label).border = thin_border
            ws.cell(row=row, column=10, value=ticket.created_at.strftime('%d/%m/%Y %H:%M')).border = thin_border
            ws.cell(row=row, column=11, value=ticket.resolved_at.strftime('%d/%m/%Y %H:%M') if ticket.resolved_at else '-').border = thin_border

            # Color code urgency
            urg_colors = {'rendah': '27AE60', 'sedang': '2E86DE', 'tinggi': 'E67E22', 'kritis': 'E74C3C'}
            urg_cell = ws.cell(row=row, column=8)
            color = urg_colors.get(ticket.urgency, 'CCCCCC')
            urg_cell.fill = PatternFill(start_color=color, end_color=color, fill_type='solid')
            urg_cell.font = Font(color='FFFFFF', bold=True)

        # Auto-fit column widths
        from openpyxl.utils import get_column_letter
        for col_idx in range(1, len(headers) + 1):
            max_length = 0
            col_letter = get_column_letter(col_idx)
            for row_idx in range(4, len(tickets) + 5):
                cell = ws.cell(row=row_idx, column=col_idx)
                try:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except:
                    pass
            ws.column_dimensions[col_letter].width = min(max_length + 4, 30)

        # Save to BytesIO
        output = BytesIO()
        wb.save(output)
        output.seek(0)

        filename = f'laporan_tiket_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
        return send_file(
            output,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=filename
        )
    except ImportError:
        from flask import flash
        flash('Module openpyxl belum terinstall. Jalankan: pip install openpyxl', 'danger')
        return redirect(request.referrer or '/')
